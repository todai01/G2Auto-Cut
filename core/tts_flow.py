"""TTS-процесс: одна запись целых фраз → нарезка → автопроверка по таблице →
прослушка и апрув.

Таблица («ТТС RU.xlsx»):
  • лист фраз — шапка с <phrase>, <name_phrase>, флагами <start>/<start_2>/<end>
    и колонками переменных (<sum_ru>, <date_ru>…), где 1 = переменная здесь;
  • листы переменных с тем же именем, что колонка (<sum_ru>, <date_ru>):
    каждая колонка листа — категория значений (<million_ru>, <day_ru>…).

Строки с флагами — кусочки одной фразы, разрезанной переменными:
start → [сумма] → start_2 → [дата] → end. Каждый кусочек сохраняется
отдельным файлом под своим именем в «Фразы_Переменные», обычная фраза —
в «Проверенные».

Запись не обязана резаться по паузам: фразы часто слипаются. Поэтому
распознанные слова всей записи идут сплошным потоком, и фразы таблицы
ищутся в нём по тексту; оговорки, повторы и значения переменных между
ними пропускаются. Разрез — в тишине между словами. Сохраняем только
после того, как дубль прослушан: автопроверка ошибается."""

import json
import os
import re
import shutil
import threading
import time
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import webview
from pydub import AudioSegment
from pydub.silence import detect_nonsilent

from core import recent_projects
from core.auto_check import (ASR_RATE, NEED_MODEL, NOISE_MAX_DBFS, NOISE_MIN_MS,
                             SPEECH_MIN_SCORE, PACK_GAP_SEC, _lat_to_cyr, _norm, _ratio,
                             similarity)
from core.num_words import synth_text
from core.ui_dialogs import ui_confirm
from utils.wav_io import export_like, fix_wav_header, wav_concat, wav_layout, wav_slice

TTS_STATE_FILE = 'tts_project.json'
TTS_STATE_VERSION = 3     # 3: фразы ищутся в сплошном потоке слов, а не по паузам
TTS_CHUNKS_DIR = 'Чанки'
TTS_DONE_DIR = 'Проверенные'
TTS_VAR_DIR = 'Фразы_Переменные'
TTS_TRASH_DIR = '_Корзина'
TTS_VALUES_DIR = 'Значения'   # готовые значения переменных: Значения/<лист>/<категория>/<значение>.wav
TTS_CUT_PAD_MS = 150
TTS_OK_SCORE = 0.80       # дубль «зелёный» — совпал с текстом
TTS_OK_MARGIN = 0.10      # …и заметно лучше второго варианта
TTS_TIE = 0.03            # почти равные варианты — решает порядок записи
TTS_LISTEN_SHARE = 0.8    # «прослушан» = отыграл хотя бы 80% длины
OUT_RATE = 8000           # готовые файлы — 8 кГц, 16 бит, моно (формат IVR)
OUT_SPEECH_DB = -18.0     # уровень речи готовых файлов: одинаковый у всех, «не тихо и не громко»
OUT_PEAK_DB = -1.0        # пики не выше — без перегруза
OUT_GAIN_LIMIT_DB = 30.0  # больше не поднимаем (если это тишина/шум)
EDIT_CONTEXT_MS = 1500    # запас записи вокруг дубля в «Сборке» — край можно вытянуть
EDIT_BUCKET_MS = 10       # шаг волны

FLAG_COLS = ('start', 'start_2', 'end')
BAD_NAME = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def _tag(v):
    """'<sum_ru>' → 'sum_ru'; не тег — None."""
    s = str(v or '').strip()
    m = re.fullmatch(r'<\s*([^<>]+?)\s*>', s)
    return m.group(1).strip().lower() if m else None


def _flag(v):
    return v not in (None, '', 0, '0') and str(v).strip() not in ('', '0')


def _lang_of(name):
    low = (name or '').lower()
    if re.search(r'(^|[_\s-])(kz|kk|kaz)([_\s>-]|$)', low):
        return 'kz'
    return 'ru'


def var_label(key):
    k = (key or '').lower()
    base = re.sub(r'_(ru|kz|kk)$', '', k)
    names = {'sum': 'Сумма', 'date': 'Дата', 'million': 'Миллионы', 'hundred_thousand': 'Сотни тысяч',
             'thousand': 'Тысячи', 'hundred': 'Сотни', 'tenge': 'Тенге', 'day': 'День', 'month': 'Месяц',
             'year': 'Год', 'name': 'Имя', 'time': 'Время'}
    label = names.get(base, base or k)
    lang = re.search(r'_(ru|kz|kk)$', k)
    return f"{label} {lang.group(1).upper()}" if lang and label != k else label


def clean_name(text):
    name = BAD_NAME.sub(' ', str(text)).strip()
    return re.sub(r'\s+', ' ', name).rstrip('. ')[:120]


def _common_name(names):
    """ru_da_1_1, ru_da_1_2, ru_da_1_3 → ru_da_1."""
    names = [n for n in names if n]
    if not names:
        return ''
    if len(names) == 1:
        return names[0]
    pref = os.path.commonprefix(names)
    pref = re.sub(r'[_\-\s]+\d*$', '', pref) if not pref.endswith(('_', '-', ' ')) else pref.rstrip('_- ')
    return pref or names[0]


def _text_tag(t):
    """<phrase>, <full_phrase_ru>, <phrase_kz> — колонка с текстом фраз."""
    return bool(t) and re.fullmatch(r'(full_)?phrase(_(ru|kz|kk))?', t) is not None


def _name_tag(t):
    """<name_phrase>, <name_phrase_ru>, <name_full_phrase_kz> — колонка с именем файла."""
    return bool(t) and re.fullmatch(r'name_(full_)?phrase(_(ru|kz|kk))?', t) is not None


def parse_tts_workbook(path):
    """Разбирает таблицу. {phrase_sheets: [...], var_sheets: {...}}."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets, var_sheets = [], {}
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        head_idx = None
        for i, r in enumerate(rows[:10]):
            if any(_text_tag(_tag(c)) for c in r):
                head_idx = i
                break
        if head_idx is not None:
            sheets.append(_parse_phrase_sheet(ws.title, rows, head_idx))
            continue
        head = rows[0]
        cols = [(c, _tag(h)) for c, h in enumerate(head) if _tag(h)]
        if not cols:
            continue
        cats = []
        for c, key in cols:
            seen, values = set(), []
            for r in rows[1:]:
                v = r[c] if c < len(r) else None
                if v is None or str(v).strip() == '':
                    continue
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                s = str(v).strip()
                if s not in seen:          # повторы-«заглушки» (100 после 900) не нужны
                    seen.add(s)
                    values.append(s)
            cats.append({"key": key, "label": var_label(key), "values": values})
        var_sheets[_tag(ws.title) or ws.title.strip().lower()] = {
            "title": ws.title, "label": var_label(_tag(ws.title) or ws.title), "categories": cats}
    wb.close()
    return {"phrase_sheets": sheets, "var_sheets": var_sheets}


def _parse_phrase_sheet(title, rows, head_idx):
    head = rows[head_idx]
    col = {}
    for c, h in enumerate(head):
        t = _tag(h)
        if t:
            col.setdefault(t, c)
    text_t = next((t for t in col if _text_tag(t)), None)
    name_t = next((t for t in col if _name_tag(t)), None)
    text_c, name_c = col.get(text_t), col.get(name_t)
    var_cols = [(t, c) for t, c in col.items() if t not in (text_t, name_t) + FLAG_COLS]
    lang_m = re.search(r'_(ru|kz|kk)$', text_t or '')
    lang = ('kz' if lang_m.group(1) in ('kz', 'kk') else 'ru') if lang_m else _lang_of(title)

    def cell(r, c):
        return r[c] if c is not None and c < len(r) else None

    lines = []
    for i, r in enumerate(rows[head_idx + 1:], start=head_idx + 2):
        text = str(cell(r, text_c) or '').strip()
        name = clean_name(cell(r, name_c) or '')
        role = next((f for f in FLAG_COLS if _flag(cell(r, col.get(f)))), None)
        vars_ = [t for t, c in var_cols if _flag(cell(r, c))]
        if not text and not name and not role:
            continue
        lines.append({"row": i, "text": text, "name": name or f"фраза_{i:04d}", "role": role, "vars": vars_})
    units, chains = build_units(lines)
    # Колонки разметки (1-based, как в Excel) — чтобы дописывать отметки в таблицу.
    mark_cols = {t: c + 1 for t, c in col.items() if t in FLAG_COLS or t in dict(var_cols)}
    return {"title": title, "lang": lang, "units": units, "chains": chains, "lines": lines,
            "var_keys": [t for t, _ in var_cols], "mark_cols": mark_cols, "head_row": head_idx + 1}


def build_units(lines):
    """Строки таблицы → юниты и цепочки.
    Юнит — каждая строка: сохраняется отдельным файлом под своим именем.
    Строки с флагами связываем в «цепочку» start → start_2 → end — это
    кусочки одной фразы. Цепочка нужна для «Сборки», суфлёра и для разрезки
    дубля, где фраза прочитана целиком (с примером значения внутри)."""
    units, chains, cur = [], [], None

    def close():
        nonlocal cur
        if cur:
            chains.append(cur)
        cur = None

    for ln in lines:
        uid = len(units)
        units.append({"id": uid, "kind": "piece" if ln["role"] else "phrase", "name": ln["name"],
                      "text": ln["text"] or ln["name"], "role": ln["role"], "row": ln["row"],
                      "vars": list(ln["vars"]), "var_before": None, "var_after": None, "chain": None})
        if ln["role"] == 'start' or (ln["role"] and cur is None):
            close()
            cur = {"id": len(chains), "parts": [uid]}
        elif ln["role"]:
            cur["parts"].append(uid)
        else:
            close()
        if ln["role"] == 'end':
            close()
    close()

    # У start/start_2 отмеченная переменная идёт после кусочка, у end — перед ним.
    for ch in chains:
        parts = [units[i] for i in ch["parts"]]
        for k, u in enumerate(parts):
            u["chain"] = ch["id"]
            if u["role"] in ('start', 'start_2'):
                nxt = parts[k + 1] if k + 1 < len(parts) else None
                var = u["vars"][0] if u["vars"] else (nxt["vars"][0] if nxt and nxt["vars"] else None)
                u["var_after"] = var
                if nxt:
                    nxt["var_before"] = var
            elif u["role"] == 'end' and not u["var_before"] and u["vars"]:
                u["var_before"] = u["vars"][0]
        ch["name"] = _common_name([u["name"] for u in parts])
        ch["vars"] = [u["var_after"] for u in parts if u["var_after"]]
    return units, chains


FIND_SCORE = 0.75         # окно слов засчитываем как фразу из таблицы
FIND_SCORE_SHORT = 0.90   # …а для фраз из 1–2 слов — только почти дословно
LONGER_TIE = 0.05         # почти равные окна — берём то, что длиннее (целая фраза, а не её начало)
LEAD_S = 0.40             # сколько тишины оставлять перед/после фразы (макс.)


def end_word_ok(window_words, target):
    """Фраза из 3+ слов засчитывается, только если сказан её конец:
    «вам сейчас удобно…» — ещё не «Вам сейчас удобно говорить?»."""
    tw = target.split()
    if len(tw) < 3 or not window_words:
        return True
    last = window_words[-1]
    return last[:4] == tw[-1][:4] or _ratio(last, tw[-1]) >= 0.6


def _first_word_ok(word, first):
    return bool(word) and bool(first) and (word[:3] == first[:3] or _ratio(word, first) >= 0.6)


def align_stream(words, units, expected=0):
    """Ищет фразы таблицы в сплошном потоке распознанных слов — не важно,
    как запись разрезалась по паузам (фразы часто слипаются в один кусок).
    words: [(слово, начало_с, конец_с)] по всей записи.
    → [{unit, score, status, i, j}] по порядку: слова [i, j) — эта фраза.
    Слова между находками (оговорки, «нет», значения переменных) пропускаются."""
    wn = [_lat_to_cyr(_norm(w[0])) for w in words]
    targets = []
    for u in units:
        t = _lat_to_cyr(_norm(u["text"]))
        if t:
            targets.append((u["id"], t, len(t.split()), t.split()[0]))
    found, i = [], 0
    while i < len(wn):
        cands = []
        for uid, t, n, first in targets:
            if not _first_word_ok(wn[i], first):
                continue
            best = (0.0, i + 1)
            for size in range(max(1, n - 2), n + 3):
                j = i + size
                if j > len(wn):
                    break
                if not end_word_ok(wn[i:j], t):
                    continue
                sc = _ratio(' '.join(wn[i:j]), t)
                if sc > best[0] + 1e-9:
                    best = (sc, j)
            need = FIND_SCORE_SHORT if n <= 2 else FIND_SCORE
            if best[0] >= need:
                cands.append((best[0], best[1], uid, n))
        if not cands:
            i += 1
            continue
        top = max(c[0] for c in cands)
        near = [c for c in cands if top - c[0] <= LONGER_TIE]
        longest = max(c[3] for c in near)
        near = [c for c in near if c[3] == longest]
        # Одинаковые тексты в разных строках — решает порядок чтения таблицы.
        sc, j, uid, _ = min(near, key=lambda c: (-round(c[0], 2), c[2] < expected, abs(c[2] - expected)))
        found.append({"unit": uid, "score": round(sc, 3), "i": i, "j": j,
                      "status": "ok" if sc >= 0.85 else "doubt"})
        expected = uid + 1
        i = j
    return found


class TtsFlowMixin:

    # ---------------- состояние ----------------

    def _tts(self):
        st = getattr(self, 'tts', None)
        if st is None:
            st = self.tts = {"v": TTS_STATE_VERSION, "excel_path": None, "sheet": None, "lang": "ru",
                             "units": [], "chains": [], "var_sheets": {},
                             "work_dir": None, "raw_path": None, "takes": [], "index": 0}
        return st

    def _tts_save(self):
        st = self._tts()
        wd = st.get("work_dir")
        if not wd or not os.path.isdir(wd):
            return
        try:
            with open(os.path.join(wd, TTS_STATE_FILE), 'w', encoding='utf-8') as f:
                json.dump(st, f, ensure_ascii=False, indent=1)
        except OSError:
            pass
        recent_projects.touch(wd, os.path.basename(wd))

    def _tts_push(self, payload):
        try:
            webview.windows[0].evaluate_js(f"ttsProgress({json.dumps(payload, ensure_ascii=False)})")
        except Exception:
            pass

    def _tts_target(self, unit):
        folder = TTS_VAR_DIR if unit["kind"] == "piece" else TTS_DONE_DIR
        return os.path.join(self._tts()["work_dir"], folder, f"{unit['name']}.wav")

    def tts_state(self):
        st = self._tts()
        units = st["units"]
        wd = st.get("work_dir")
        saved = {u["id"]: bool(wd) and os.path.exists(self._tts_target(u)) for u in units}
        takes_by_unit = {}
        for i, t in enumerate(st["takes"]):
            if t.get("unit") is not None:
                takes_by_unit.setdefault(t["unit"], []).append(i)
        takes = []
        for i, t in enumerate(st["takes"]):
            u = units[t["unit"]] if t.get("unit") is not None and t["unit"] < len(units) else None
            same = takes_by_unit.get(t.get("unit"), [])
            takes.append({**t, "i": i, "name": u["name"] if u else None, "kind": u["kind"] if u else None,
                          "unit_saved": saved.get(t.get("unit"), False),
                          "take_no": same.index(i) + 1 if i in same else 0, "take_total": len(same)})
        pieces = [u for u in units if u["kind"] == "piece"]
        return {
            "loaded": bool(units), "has_takes": bool(st["takes"]),
            "excel_name": os.path.basename(st["excel_path"] or ''), "sheet": st.get("sheet"),
            "lang": st.get("lang", "ru"), "project": os.path.basename(wd) if wd else '',
            "units": [{**u, "saved": saved[u["id"]], "takes": takes_by_unit.get(u["id"], [])} for u in units],
            "chains": st.get("chains", []),
            "var_cols": [{"key": k, "label": var_label(k)} for k in (st.get("var_cols") or [])],
            "marks_pending": len(st.get("marks_dirty") or []),
            "takes": takes, "index": min(st.get("index", 0), max(0, len(takes) - 1)),
            "var_sheets": {k: {"label": v["label"],
                               "categories": [{"key": c["key"], "label": c["label"], "count": len(c["values"])}
                                              for c in v["categories"]]}
                           for k, v in st["var_sheets"].items()},
            "stats": {"units": len(units), "saved": sum(saved.values()),
                      "pieces": len(pieces), "pieces_saved": sum(1 for u in pieces if saved[u["id"]]),
                      "chains": len(st.get("chains", [])),
                      "missing": sum(1 for u in units if not saved[u["id"]] and u["id"] not in takes_by_unit),
                      "doubt": sum(1 for t in st["takes"] if t.get("status") == "doubt"
                                   and not saved.get(t.get("unit"), False))},
        }

    # ---------------- загрузка ----------------

    def tts_pick_excel(self):
        picked = webview.windows[0].create_file_dialog(
            webview.FileDialog.OPEN, file_types=('Excel files (*.xlsx)', 'All files (*.*)'))
        if not picked:
            return {"error": "cancel"}
        path = picked if isinstance(picked, str) else picked[0]
        return self.tts_load_excel(path)

    def tts_load_excel(self, path, sheet=None):
        try:
            parsed = parse_tts_workbook(path)
        except Exception as e:
            return {"error": f"Не удалось прочитать таблицу «{os.path.basename(path)}».\n\n{e}"}
        sheets = parsed["phrase_sheets"]
        if not sheets:
            return {"error": "В таблице нет листа фраз: нужна шапка с колонкой <phrase> (или <full_phrase_ru>)."}
        if sheet is None and len(sheets) > 1:
            self._tts_pending = (path, parsed)
            return {"choose_sheet": [{"title": s["title"], "lang": s["lang"], "units": len(s["units"])}
                                     for s in sheets]}
        chosen = next((s for s in sheets if s["title"] == sheet), sheets[0])
        st = self._tts()
        if st["takes"] and st.get("excel_path") and os.path.abspath(st["excel_path"]) != os.path.abspath(path):
            st["takes"] = []      # другая таблица — старая привязка дублей неверна
        st.update({"excel_path": path, "sheet": chosen["title"], "lang": chosen["lang"],
                   "units": chosen["units"], "chains": chosen["chains"], "v": TTS_STATE_VERSION,
                   "lines": chosen["lines"], "mark_cols": chosen["mark_cols"], "var_cols": chosen["var_keys"],
                   "head_row": chosen["head_row"], "marks_dirty": [],
                   "var_sheets": {k: v for k, v in parsed["var_sheets"].items()
                                  if k in chosen["var_keys"] or not chosen["var_keys"]}})
        st["index"] = 0
        self._tts_save()
        return self.tts_state()

    def tts_choose_sheet(self, title):
        pend = getattr(self, '_tts_pending', None)
        if not pend:
            return {"error": "Таблица не выбрана."}
        return self.tts_load_excel(pend[0], title)

    def tts_open_project(self, path):
        """Продолжить TTS-проект (tts_project.json в папке)."""
        try:
            with open(os.path.join(path, TTS_STATE_FILE), 'r', encoding='utf-8') as f:
                data = json.load(f)
        except (OSError, ValueError):
            return {"error": "Не нашёл сохранённый TTS-проект в этой папке."}
        self.tts = None
        st = self._tts()
        old = data.get("v") != TTS_STATE_VERSION
        st.update(data)
        st["work_dir"] = path
        # Запись могла оборваться (закрыли программу во время записи).
        for raw in {t.get("raw") for t in st["takes"]} | {st.get("raw_path")}:
            if raw and os.path.exists(raw):
                fix_wav_header(raw)
        if old:
            # Проект старого формата (цепочка = один файл): перечитываем
            # таблицу, а дубли нужно нарезать заново — привязка другая.
            st["takes"], st["index"] = [], 0
            if st.get("excel_path") and os.path.exists(st["excel_path"]):
                res = self.tts_load_excel(st["excel_path"], st.get("sheet"))
                if res.get("error"):
                    return res
            st["v"] = TTS_STATE_VERSION
            self._tts_save()
        return self.tts_state()

    @staticmethod
    def tts_is_project(path):
        return bool(path) and os.path.exists(os.path.join(path, TTS_STATE_FILE))

    def tts_pick_recording(self):
        st = self._tts()
        if not st["units"]:
            return {"error": "Сначала загрузите таблицу."}
        picked = webview.windows[0].create_file_dialog(
            webview.FileDialog.OPEN, file_types=('Audio Files (*.wav;*.mp3)', 'All files (*.*)'))
        if not picked:
            return {"error": "cancel"}
        path = picked if isinstance(picked, str) else picked[0]
        if getattr(self, '_tts_rec', None):
            return {"error": "Сейчас идёт запись с микрофона — сначала остановите её."}
        if st["takes"] and not ui_confirm("Запись уже нарезана и проверена. Нарезать заново?<br><br>"
                                          "Сохранённые фразы останутся на месте.", "Нарезать заново", "Отмена"):
            return {"error": "cancel"}
        threading.Thread(target=self._tts_process, args=(path,), daemon=True).start()
        return {"status": "started"}

    # ---------------- нарезка + распознавание ----------------

    def _tts_process(self, path):
        lock = getattr(self, '_auto_check_lock', None)
        if lock is None:
            lock = self._auto_check_lock = threading.Lock()
        if not lock.acquire(blocking=False):
            self._tts_push({"stage": "error", "error": "Распознавание уже идёт."})
            return
        try:
            self._auto_check_stop = False
            res = self._tts_cut_and_check(path)
        except Exception as e:
            res = {"error": f"Не удалось обработать запись: {e}"}
        finally:
            lock.release()
        if res.get("error"):
            self._tts_push({"stage": "error", **res})
        else:
            self._tts_push({"stage": "done", "state": self.tts_state()})

    def _tts_cut_and_check(self, path):
        st = self._tts()
        self._tts_push({"stage": "cut", "done": 0, "total": 0})
        audio = AudioSegment.from_file(path)
        levels = self._frame_levels(audio)
        if not levels:
            return {"error": "Запись пустая или слишком короткая."}
        _, _, _, thresh, pause, _ = self._suggest_cut_settings(levels, len(st["units"]))
        ranges = detect_nonsilent(audio, min_silence_len=int(pause), silence_thresh=int(thresh))
        if not ranges:
            return {"error": "В записи не нашлось речи."}

        work_dir = os.path.dirname(path)
        chunks_dir = os.path.join(work_dir, TTS_CHUNKS_DIR)
        os.makedirs(chunks_dir, exist_ok=True)
        for f in os.listdir(chunks_dir):
            if re.fullmatch(r'\d{4}(_\d+)?\.wav', f):
                try:
                    os.remove(os.path.join(chunks_dir, f))
                except OSError:
                    pass
        pieces = []
        for i, (a, b) in enumerate(ranges):
            piece = audio[max(0, a - TTS_CUT_PAD_MS):min(len(audio), b + TTS_CUT_PAD_MS)]
            pieces.append((f"{i + 1:04d}", piece, a))
            if i % 20 == 0:
                self._tts_push({"stage": "cut", "done": i + 1, "total": len(ranges)})

        self._tts_push({"stage": "model", "done": 0, "total": 0})
        model, err = self._auto_check_model()
        if err is NEED_MODEL:
            return {"error": "need_model"}
        if err:
            return {"error": err}

        heards, words = self._tts_transcribe([p[1] for p in pieces], 'kk' if st["lang"] == 'kz' else 'ru', model)
        takes = self._tts_build_takes(audio, pieces, words, st["units"], chunks_dir, path)
        st.update({"work_dir": work_dir, "raw_path": path, "index": 0, "takes": takes})
        st["index"] = self._tts_next(-1)
        self._tts_save()
        return {"status": "ok"}

    def _tts_build_takes(self, audio, pieces, words, units, chunks_dir, raw_path):
        """Сплошной поток слов всей записи → фразы таблицы → по файлу на
        каждую найденную фразу (разрез — в тишине между словами)."""
        stream = []
        for k, (name, piece, a) in enumerate(pieces):
            off = max(0, a - TTS_CUT_PAD_MS) / 1000.0
            stream += [(w, off + ws, off + we, k) for w, ws, we in words[k]]
        spans = self._tts_spans(stream, align_stream(stream, units))
        return self._tts_cut_spans(audio, 0, stream, spans, raw_path, chunks_dir, 1)

    @staticmethod
    def _tts_spans(stream, found):
        """Найденные фразы + неопознанная речь (≥4 слов подряд — вдруг это
        фраза, которую распознавание исказило; привяжете вручную)."""
        spans = []
        used = [False] * len(stream)
        for f in found:
            i, j = f["i"], f["j"]
            for x in range(i, j):
                used[x] = True
            spans.append((i, j, {"unit": f["unit"], "score": f["score"], "status": f["status"]}))
        x = 0
        while x < len(stream):
            if used[x]:
                x += 1
                continue
            y = x
            while y < len(stream) and not used[y] and stream[y][3] == stream[x][3]:
                y += 1
            if y - x >= 4:
                spans.append((x, y, {"unit": None, "score": 0.0, "status": "none"}))
            x = y
        spans.sort(key=lambda sp: sp[0])
        return spans

    def _tts_cut_spans(self, audio, offset_ms, stream, spans, raw_path, chunks_dir, first_n, push=True):
        """Вырезает каждую фразу из исходника байт в байт. audio — анализный
        звук, начинающийся с offset_ms записи (для поиска тишины); время слов
        в stream — от начала записи."""
        takes = []
        end_s = (offset_ms + len(audio)) / 1000.0
        rel = lambda sec: sec - offset_ms / 1000.0
        for n, (i, j, info) in enumerate(spans):
            first_s, last_e = stream[i][1], stream[j - 1][2]
            prev_e = stream[i - 1][2] if i > 0 else offset_ms / 1000.0
            next_s = stream[j][1] if j < len(stream) else end_s
            lo = offset_ms + self._tts_quiet_point(audio, *self._tts_edge_window(
                rel(max(prev_e, first_s - LEAD_S)), rel(first_s + 0.03)))
            hi = offset_ms + self._tts_quiet_point(audio, *self._tts_edge_window(
                rel(last_e - 0.03), rel(min(next_s, last_e + LEAD_S))))
            if hi - lo < 150:
                lo, hi = int(first_s * 1000), int(last_e * 1000) + 60
            name = f"{first_n + n:04d}.wav"
            # Копия байт из исходника — без перекодирования, качество не трогаем.
            if not wav_slice(raw_path, os.path.join(chunks_dir, name), lo, hi):
                audio[lo - offset_ms:hi - offset_ms].export(os.path.join(chunks_dir, name), format="wav")
            takes.append({"file": os.path.join(TTS_CHUNKS_DIR, name), "start_ms": lo, "ms": hi - lo, "orig": [lo, hi],
                          "raw": raw_path, "heard": ''.join(w[0] for w in stream[i:j]).strip(), "listened": False,
                          "chunk": first_n + n, **info})
            if push and n % 20 == 0:
                self._tts_push({"stage": "save", "done": n + 1, "total": len(spans)})
        return takes

    @staticmethod
    def _tts_edge_window(a, b):
        """Время слов у распознавания неточное (±50 мс): паузу между словами
        в 50–100 мс оно может «съесть». Узкое окно расширяем вокруг середины."""
        if b - a < 0.16:
            c = (a + b) / 2
            return c - 0.10, c + 0.10
        return a, b

    @staticmethod
    def _tts_quiet_point(seg, lo_s, hi_s):
        """Самое тихое место (мс) в окне — туда и ставим разрез."""
        lo, hi = max(0, int(lo_s * 1000)), min(len(seg), int(hi_s * 1000))
        if hi - lo < 20:
            return max(0, min(len(seg), int((lo_s + hi_s) * 500)))
        frames = [(t, seg[t:t + 10].rms) for t in range(lo, hi - 10 + 1, 5)]
        floor = min(r for _, r in frames)
        quiet = floor * 1.5 + 30
        # Середина самой длинной тихой полосы — не край паузы, а её центр.
        best, run = (0, frames[0][0]), None
        for t, r in frames:
            if r <= quiet:
                run = run or t
                if t - run >= best[0]:
                    best = (t - run, run)
            else:
                run = None
        return best[1] + (best[0] + 10) // 2

    def _tts_transcribe_pack(self, model, pack, language, live=False, hotwords=None):
        """Как _asr_transcribe_pack, но со временем слов внутри каждого дубля:
        [(слово, начало_с, конец_с)] — по ним режем цепочку на кусочки.
        live — запись с микрофона: чувствительнее к тихим коротким словам
        (команды), hotwords — подсказка модели, какие слова ждать."""
        import numpy as np
        gap = np.zeros(int(PACK_GAP_SEC * ASR_RATE), dtype=np.float32)
        parts, spans, t = [], [], 0.0
        for k, it in enumerate(pack):
            if k:
                parts.append(gap)
                t += PACK_GAP_SEC
            parts.append(it['audio'])
            d = len(it['audio']) / ASR_RATE
            spans.append((t, t + d))
            t += d
        audio = np.concatenate(parts) if len(parts) > 1 else parts[0]
        # Тихая запись (некоторые гарнитуры) распознаётся хуже — поднимаем по
        # уровню голоса, а не по одному пику (щелчок не должен мешать).
        if len(audio):
            ref = float(np.percentile(np.abs(audio), 99.5))
            if 1e-4 < ref < 0.35:
                audio = np.clip(audio * min(40.0, 0.5 / ref), -1.0, 1.0)
        opts = dict(language=None if language == 'auto' else language, beam_size=1, best_of=1,
                    word_timestamps=True, vad_filter=True, condition_on_previous_text=False)
        if live:
            opts["vad_parameters"] = {"threshold": 0.3, "min_speech_duration_ms": 80, "min_silence_duration_ms": 300}
        if hotwords:
            opts["hotwords"] = hotwords
        try:
            segments, _info = model.transcribe(audio, **opts)
        except TypeError:                          # старая faster-whisper — без hotwords
            opts.pop("hotwords", None)
            segments, _info = model.transcribe(audio, **opts)
        no_speech, logprob = (0.85, -1.0) if live else (0.6, -0.7)
        words = [[] for _ in pack]
        for seg in segments:
            if getattr(seg, 'no_speech_prob', 0) > no_speech and getattr(seg, 'avg_logprob', 0) < logprob:
                continue
            for w in (seg.words or []):
                mid = (w.start + w.end) / 2
                for k, (a, b) in enumerate(spans):
                    if a - 0.3 <= mid <= b + 0.3:
                        words[k].append((w.word, max(0.0, w.start - a), max(0.0, w.end - a)))
                        break
        return words

    def _tts_transcribe(self, segs, language, model):
        """[текст дубля], [[(слово, начало, конец)]]."""
        total = len(segs)
        heard = [''] * total
        words_out = [[] for _ in range(total)]
        state = {"done": 0}
        lock = threading.Lock()

        def finish(k, text, words=None):
            with lock:
                heard[k] = text
                words_out[k] = words or []
                state["done"] += 1
                d = state["done"]
            if d % 5 == 0 or d == total:
                self._tts_push({"stage": "asr", "done": d, "total": total})

        items = []
        for k, seg in enumerate(segs):
            samples, ms, peak = self._audit_asr_samples(seg)
            if ms < NOISE_MIN_MS or peak < NOISE_MAX_DBFS:
                finish(k, '')
            else:
                items.append({"index": k, "audio": samples, "ms": ms})

        def run(pack):
            if getattr(self, '_auto_check_stop', False):
                return
            try:
                out = self._tts_transcribe_pack(model, pack, language)
            except Exception:
                out = [[] for _ in pack]
            for it, ws in zip(pack, out):
                h = ''.join(w[0] for w in ws).strip()
                if self._asr_is_noise(h):
                    finish(it['index'], '')
                else:
                    finish(it['index'], h, ws)

        with ThreadPoolExecutor(max_workers=getattr(self, '_asr_workers', 1)) as ex:
            list(ex.map(run, self._asr_packs(items)))
        return heard, words_out

    def tts_stop(self):
        self._auto_check_stop = True
        return {"status": "ok"}

    # ---------------- прослушка и апрув ----------------

    def _tts_mark_listened(self):
        p = getattr(self, '_tts_playing', None)
        if not p:
            return
        i, started, dur = p
        takes = self._tts()["takes"]
        if i < len(takes) and time.time() - started >= dur * TTS_LISTEN_SHARE:
            takes[i]["listened"] = True
        self._tts_playing = None

    def tts_play(self, i):
        st = self._tts()
        self._tts_mark_listened()
        if not (0 <= i < len(st["takes"])):
            return {"playing": False}
        path = os.path.join(st["work_dir"], st["takes"][i]["file"])
        if not os.path.exists(path):
            return {"playing": False, "error": "Файл дубля пропал с диска."}
        st["index"] = i
        dur = st["takes"][i].get("ms", 0) / 1000.0
        # Слушаете ровно то, что сохранится: громкость выровнена, 8 кГц.
        try:
            path = self._tts_preview(path)
        except Exception:
            pass
        self.player.play(path)
        self._tts_playing = (i, time.time(), dur)
        return {"playing": True, "duration": dur}

    def tts_stop_audio(self):
        self._tts_mark_listened()
        self.player.stop()
        return {"playing": False}

    def tts_select(self, i):
        st = self._tts()
        self._tts_mark_listened()
        if 0 <= i < len(st["takes"]):
            st["index"] = i
        return self.tts_state()

    def _tts_next(self, after):
        """Следующий дубль, который ещё нужно прослушать: привязан к фразе,
        а фраза не сохранена. Сначала — вперёд по записи, потом с начала.
        Из повторов одной фразы предлагаем последний (диктор переговаривает
        после ошибки); забракуете его — следующим пойдёт предыдущий."""
        st = self._tts()
        takes = st["takes"]
        n = len(takes)
        last = {}
        for k, t in enumerate(takes):
            if t.get("unit") is not None and not t.get("rejected"):
                last[t["unit"]] = k
        saved = {}
        for k in list(range(after + 1, n)) + list(range(0, after + 1)):
            t = takes[k]
            uid = t.get("unit")
            if uid is None or t.get("rejected") or last.get(uid) != k:
                continue
            if uid not in saved:
                saved[uid] = os.path.exists(self._tts_target(st["units"][uid])) if st.get("work_dir") else False
            if not saved[uid]:
                return k
        return max(0, min(after, n - 1))

    def tts_approve(self, i):
        st = self._tts()
        self._tts_mark_listened()
        if not (0 <= i < len(st["takes"])):
            return {"error": "Нет такого дубля."}
        t = st["takes"][i]
        if t.get("unit") is None:
            return {"error": "Дубль не привязан к фразе — выберите фразу."}
        if not t.get("listened"):
            return {"need_listen": True}
        unit = st["units"][t["unit"]]
        target = self._tts_target(unit)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        if os.path.exists(target):
            self._tts_trash(target)
        self._tts_render_final(os.path.join(st["work_dir"], t["file"]), target)
        for other in st["takes"]:
            if other.get("unit") == t["unit"]:
                other["approved"] = False
        t["approved"] = True
        t["rejected"] = False
        st["index"] = self._tts_next(i)
        self._tts_save()
        return {**self.tts_state(), "saved_kind": unit["kind"], "saved_name": unit["name"]}

    def tts_reject(self, i):
        """«Плохой дубль» — убрать из очереди (файл чанка остаётся)."""
        st = self._tts()
        if 0 <= i < len(st["takes"]):
            st["takes"][i]["rejected"] = not st["takes"][i].get("rejected")
            if st["takes"][i]["rejected"]:
                # Забраковали попытку — сразу предлагаем предыдущую попытку той же фразы.
                uid = st["takes"][i].get("unit")
                prev = [k for k in range(i) if uid is not None and st["takes"][k].get("unit") == uid
                        and not st["takes"][k].get("rejected")]
                st["index"] = prev[-1] if prev else self._tts_next(i)
            self._tts_save()
        return self.tts_state()

    def tts_unapprove(self, unit_id):
        st = self._tts()
        if not (0 <= unit_id < len(st["units"])):
            return self.tts_state()
        target = self._tts_target(st["units"][unit_id])
        if os.path.exists(target):
            self._tts_trash(target)
        for t in st["takes"]:
            if t.get("unit") == unit_id:
                t["approved"] = False
        self._tts_save()
        return self.tts_state()

    def tts_assign(self, i, unit_id):
        """Ручная привязка дубля к другой фразе (если автопроверка ошиблась)."""
        st = self._tts()
        if 0 <= i < len(st["takes"]):
            t = st["takes"][i]
            if unit_id is None or unit_id < 0:
                t.update({"unit": None, "status": "none", "score": 0.0})
            elif unit_id < len(st["units"]):
                score = similarity(t.get("heard") or '', st["units"][unit_id]["text"]) if t.get("heard") else 0.0
                t.update({"unit": unit_id, "score": round(score, 3), "status": "manual"})
            t["approved"] = False
            self._tts_save()
        return self.tts_state()

    def tts_play_unit(self, unit_id):
        """Прослушать уже сохранённую фразу (карточка «Переменные», список)."""
        st = self._tts()
        if not (0 <= unit_id < len(st["units"])):
            return {"playing": False}
        target = self._tts_target(st["units"][unit_id])
        if not os.path.exists(target):
            return {"playing": False}
        self._tts_mark_listened()
        self.player.play(target)
        return {"playing": True, "duration": len(AudioSegment.from_file(target)) / 1000.0}

    def _tts_trash(self, path):
        trash = os.path.join(self._tts()["work_dir"], TTS_TRASH_DIR)
        os.makedirs(trash, exist_ok=True)
        base, ext = os.path.splitext(os.path.basename(path))
        dest = os.path.join(trash, f"{base}{ext}")
        n = 2
        while os.path.exists(dest):
            dest = os.path.join(trash, f"{base} ({n}){ext}")
            n += 1
        shutil.move(path, dest)

    def tts_open_folder(self, which):
        st = self._tts()
        if not st.get("work_dir"):
            return {"error": "Проект ещё не создан."}
        folder = os.path.join(st["work_dir"], TTS_VAR_DIR if which == 'var' else TTS_DONE_DIR)
        os.makedirs(folder, exist_ok=True)
        try:
            os.startfile(folder)
        except Exception as e:
            return {"error": str(e)}
        return {"status": "ok"}

    # ---------------- «Сборка»: монтаж дубля прямо в программе ----------------
    # Дубль показывается вместе с запасом записи вокруг (EDIT_CONTEXT_MS):
    # если автонарезка срезала слово, край можно просто вытянуть наружу.
    # Все правки собираются из исходной записи байт в байт.

    def _tts_raw(self, t):
        """Исходная запись дубля (у живой записи их может быть несколько)."""
        return (t or {}).get("raw") or self._tts().get("raw_path")

    @staticmethod
    def _tts_orig(t):
        """Границы дубля, как его нарезала программа (до правок)."""
        if not t.get("orig"):
            t["orig"] = [t["start_ms"], t["start_ms"] + t["ms"]]
        return t["orig"]

    def _tts_take_ranges(self, t, a=None, b=None, cuts=None):
        """Абсолютные куски исходника (мс), из которых состоит дубль."""
        ed = t.get("edit") or {}
        orig = self._tts_orig(t)
        a = ed.get("a", orig[0]) if a is None else a
        b = ed.get("b", orig[1]) if b is None else b
        cuts = ed.get("cuts", []) if cuts is None else cuts
        ranges, cur = [], a
        for c0, c1 in sorted((max(a, c0), min(b, c1)) for c0, c1 in cuts):
            if c1 <= c0:
                continue
            if c0 > cur:
                ranges.append((cur, c0))
            cur = max(cur, c1)
        if b > cur:
            ranges.append((cur, b))
        return ranges

    def tts_editor_load(self, i):
        st = self._tts()
        t = st["takes"][i] if 0 <= i < len(st["takes"]) else None
        raw = self._tts_raw(t) if t else None
        if not raw or not os.path.exists(raw):
            return {"error": "Нет исходной записи для этого дубля."}
        ed = t.get("edit") or {}
        orig = self._tts_orig(t)
        a, b = ed.get("a", orig[0]), ed.get("b", orig[1])
        total = getattr(self, '_tts_raw_len', None)
        if not total or total[0] != raw or getattr(self, '_tts_rec', None):
            lay = wav_layout(raw)
            length = int(lay[2] / lay[3] * 1000 / lay[4]) if lay else len(AudioSegment.from_file(raw))
            total = self._tts_raw_len = (raw, length)
        lo, hi = max(0, a - EDIT_CONTEXT_MS), min(total[1], b + EDIT_CONTEXT_MS)
        import tempfile
        tmp = os.path.join(tempfile.gettempdir(), f"gvox_tts_ctx_{os.getpid()}.wav")
        if wav_slice(raw, tmp, lo, hi):
            seg = AudioSegment.from_file(tmp)
        else:
            seg = AudioSegment.from_file(raw)[lo:hi]
        seg = seg.set_channels(1).set_frame_rate(16000).set_sample_width(2)
        samples = seg.get_array_of_samples()
        per = 16000 * EDIT_BUCKET_MS // 1000
        peaks = [max(abs(min(samples[k:k + per])), max(samples[k:k + per])) / 32768
                 for k in range(0, len(samples), per)]
        # Волна в масштабе (тихая запись не должна выглядеть плоской линией).
        top = max(peaks) if peaks else 0
        k_ = min(40.0, 0.9 / top) if top > 1e-4 else 1.0
        peaks = [round(min(1.0, x * k_), 3) for x in peaks]
        return {"i": i, "lo": lo, "hi": hi, "a": a, "b": b, "cuts": ed.get("cuts", []),
                "orig": orig, "audacity": bool(t.get("audacity")), "peaks": peaks, "bucket": EDIT_BUCKET_MS}

    def tts_editor_apply(self, i, a, b, cuts):
        """Записать правку в файл дубля (из исходника, байт в байт). Правленый
        дубль нужно прослушать заново — только потом его можно сохранить."""
        st = self._tts()
        if not (0 <= i < len(st["takes"])):
            return {"error": "Нет такого дубля."}
        t = st["takes"][i]
        self._tts_orig(t)
        a, b = int(a), int(b)
        cuts = [[int(c0), int(c1)] for c0, c1 in (cuts or []) if int(c1) > int(c0)]
        if b - a < 100:
            return {"error": "Слишком короткий кусок."}
        ranges = self._tts_take_ranges(t, a, b, cuts)
        dst = os.path.join(st["work_dir"], t["file"])
        if not wav_concat(self._tts_raw(t), dst, ranges):
            src = AudioSegment.from_file(self._tts_raw(t))
            out = sum((src[s:e] for s, e in ranges[1:]), src[ranges[0][0]:ranges[0][1]])
            out.export(dst, format="wav")
        t["edit"] = {"a": a, "b": b, "cuts": cuts}
        t["ms"] = sum(e - s for s, e in ranges)
        t["listened"] = False
        t.pop("audacity", None)
        self._tts_save()
        return self.tts_state()

    def tts_editor_reset(self, i):
        """Вернуть дубль как его нарезала программа."""
        st = self._tts()
        if 0 <= i < len(st["takes"]):
            return self._tts_reset_file(i)
        return self.tts_state()

    def _tts_reset_file(self, i):
        st = self._tts()
        t = st["takes"][i]
        orig = self._tts_orig(t)
        wav_slice(self._tts_raw(t), os.path.join(st["work_dir"], t["file"]), orig[0], orig[1])
        t["ms"] = orig[1] - orig[0]
        t["listened"] = False
        t.pop("edit", None)
        t.pop("audacity", None)
        self._tts_save()
        return self.tts_state()

    # ---------------- Audacity ----------------

    def tts_send_to_audacity(self, i):
        """Дубль — в Audacity (чистый проект, один трек, масштаб по дублю)."""
        st = self._tts()
        if not (0 <= i < len(st["takes"])):
            return {"error": "Нет такого дубля."}
        if not self._ensure_audacity_ready():
            return {"error": "Audacity не отвечает. Откройте его вручную и нажмите ещё раз."}
        path = os.path.abspath(os.path.join(st["work_dir"], st["takes"][i]["file"])).replace('\\', '/')
        self.audacity.send_command('SelectAll:')
        self.audacity.send_command('RemoveTracks:')
        self.audacity.send_command(f'Import2: Filename="{path}"')
        self.audacity.send_command('SelectAll:')
        self.audacity.send_command('ZoomSel:')
        self.audacity.send_command('SelectNone:')
        self._tts_audacity_take = i
        return {"status": "ok"}

    def tts_take_from_audacity(self, i):
        """Забрать правку из Audacity в файл дубля — в формате исходной записи."""
        st = self._tts()
        if not (0 <= i < len(st["takes"])):
            return {"error": "Нет такого дубля."}
        if getattr(self, '_tts_audacity_take', None) != i:
            return {"error": "В Audacity сейчас другой дубль — сначала отправьте этот."}
        import tempfile
        tmp = os.path.join(tempfile.gettempdir(), f"gvox_tts_aud_{os.getpid()}.wav").replace('\\', '/')
        if os.path.exists(tmp):
            os.remove(tmp)
        self.audacity.send_command('SelectAll:')
        self.audacity.send_command(f'Export2: Filename="{tmp}" NumChannels=1')
        for _ in range(30):
            if os.path.exists(tmp) and os.path.getsize(tmp) > 44:
                break
            time.sleep(0.1)
        else:
            return {"error": "Audacity не сохранил файл."}
        time.sleep(0.2)
        t = st["takes"][i]
        dst = os.path.join(st["work_dir"], t["file"])
        seg = AudioSegment.from_file(tmp)
        export_like(seg, dst, self._tts_raw(t))
        t["ms"] = len(seg)
        t["listened"] = False
        t["audacity"] = True
        self._tts_save()
        return self.tts_state()

    # ---------------- экран «Переменные» ----------------
    # Значения переменных лежат в «Значения/<лист>/<категория>/<значение>.wav»
    # (лист — <sum_ru>, категория — <million_ru> и т.д.). У каждого значения —
    # текст для синтеза («пять миллионов»), его можно поправить руками.

    def _tts_value_dir(self, sheet, cat):
        return os.path.join(self._tts()["work_dir"] or '', TTS_VALUES_DIR, sheet, cat)

    def _tts_value_file(self, sheet, cat, value):
        return os.path.join(self._tts_value_dir(sheet, cat), f"{clean_name(value)}.wav")

    def _tts_value_text(self, sheet, cat, value):
        custom = (self._tts().get("var_texts") or {}).get(f"{sheet}|{cat}|{value}")
        return custom or synth_text(cat, value), bool(custom)

    def tts_vars_state(self):
        st = self._tts()
        has_dir = bool(st.get("work_dir"))
        sheets = []
        for key, sh in st["var_sheets"].items():
            cats = []
            for c in sh["categories"]:
                values = []
                for v in c["values"]:
                    text, custom = self._tts_value_text(key, c["key"], v)
                    values.append({"value": v, "text": text, "custom": custom,
                                   "ready": has_dir and os.path.exists(self._tts_value_file(key, c["key"], v))})
                cats.append({"key": c["key"], "label": c["label"], "values": values,
                             "ready": sum(1 for x in values if x["ready"])})
            sheets.append({"key": key, "label": sh["label"], "categories": cats})
        chains = []
        for ch in st.get("chains", []):
            parts = [st["units"][i] for i in ch["parts"]]
            chains.append({"id": ch["id"], "name": ch["name"], "vars": ch.get("vars", []),
                           "parts": [{"id": u["id"], "name": u["name"], "role": u["role"], "text": u["text"],
                                      "var_after": u.get("var_after"), "var_before": u.get("var_before"),
                                      "saved": has_dir and os.path.exists(self._tts_target(u))} for u in parts]})
        return {"sheets": sheets, "chains": chains, "has_project": has_dir}

    def tts_var_set_text(self, sheet, cat, value, text):
        st = self._tts()
        texts = st.setdefault("var_texts", {})
        key = f"{sheet}|{cat}|{value}"
        text = (text or '').strip()
        if text and text != synth_text(cat, value):
            texts[key] = text
        else:
            texts.pop(key, None)
        self._tts_save()
        return {"text": self._tts_value_text(sheet, cat, value)[0]}

    def tts_var_play(self, sheet, cat, value):
        path = self._tts_value_file(sheet, cat, value)
        if not os.path.exists(path):
            return {"playing": False}
        self._tts_mark_listened()
        self.player.play(path)
        return {"playing": True, "duration": len(AudioSegment.from_file(path)) / 1000.0}

    def tts_var_open_folder(self, sheet, cat):
        if not self._tts().get("work_dir"):
            return {"error": "Сначала загрузите запись — от неё зависит папка проекта."}
        folder = self._tts_value_dir(sheet, cat)
        os.makedirs(folder, exist_ok=True)
        try:
            os.startfile(folder)
        except Exception as e:
            return {"error": str(e)}
        return {"status": "ok"}

    @staticmethod
    def _tts_norm_value(text):
        return re.sub(r'\s+', ' ', str(text).lower().replace('ё', 'е')).strip()

    def tts_var_import(self, sheet, cat):
        """Разложить уже готовые файлы значений: имя файла — само значение
        («5.wav»), его текст («пять миллионов.wav») или содержит число
        («million_5.wav»). Копируем байт в байт."""
        st = self._tts()
        if not st.get("work_dir"):
            return {"error": "Сначала загрузите запись — от неё зависит папка проекта."}
        sh = st["var_sheets"].get(sheet)
        c = next((c for c in (sh or {}).get("categories", []) if c["key"] == cat), None)
        if not c:
            return {"error": "Нет такой категории."}
        picked = webview.windows[0].create_file_dialog(webview.FileDialog.FOLDER)
        if not picked:
            return {"error": "cancel"}
        folder = picked if isinstance(picked, str) else picked[0]
        by_name = {}
        for v in c["values"]:
            by_name.setdefault(self._tts_norm_value(v), v)
            by_name.setdefault(self._tts_norm_value(self._tts_value_text(sheet, cat, v)[0]), v)
        digits = {re.sub(r'\D', '', v): v for v in c["values"] if re.fullmatch(r'\d+', v)}
        matched, unknown = 0, []
        os.makedirs(self._tts_value_dir(sheet, cat), exist_ok=True)
        for f in sorted(os.listdir(folder)):
            stem, ext = os.path.splitext(f)
            if ext.lower() not in ('.wav', '.mp3'):
                continue
            v = by_name.get(self._tts_norm_value(stem))
            if v is None:
                nums = re.findall(r'\d+', stem)
                v = digits.get(nums[-1].lstrip('0') or '0') if len(nums) == 1 else None
            if v is None:
                unknown.append(f)
                continue
            src = os.path.join(folder, f)
            if ext.lower() == '.wav':
                shutil.copy2(src, self._tts_value_file(sheet, cat, v))
            else:
                AudioSegment.from_file(src).export(self._tts_value_file(sheet, cat, v), format="wav")
            matched += 1
        return {"matched": matched, "unknown": unknown[:30], "unknown_count": len(unknown),
                "state": self.tts_vars_state()}

    @staticmethod
    def _tts_speech_db(seg):
        """Громкость речи: 90-й перцентиль громкости 50-мс окон (тишина не в счёт)."""
        levels = sorted(seg[k:k + 50].dBFS for k in range(0, max(1, len(seg) - 50), 50)
                        if seg[k:k + 50].rms > 0)
        if not levels:
            return None
        return round(levels[int(len(levels) * 0.9) - 1 if len(levels) > 1 else 0], 1)

    @classmethod
    def _tts_normalized(cls, seg):
        """Выровнять громкость: речь — на OUT_SPEECH_DB, пики — не выше OUT_PEAK_DB."""
        seg = seg.set_channels(1)
        level = cls._tts_speech_db(seg)
        if level is None or seg.max_dBFS == float('-inf'):
            return seg
        gain = OUT_SPEECH_DB - level
        gain = min(gain, OUT_PEAK_DB - seg.max_dBFS, OUT_GAIN_LIMIT_DB)
        return seg.apply_gain(gain) if abs(gain) > 0.1 else seg

    @staticmethod
    def _tts_export_8k(seg, dst):
        """WAV 8 кГц / 16 бит / моно. Частоту пересчитывает ffmpeg (с фильтром,
        без призвуков); нет ffmpeg — запасной путь pydub."""
        try:
            seg.export(dst, format="wav", parameters=["-ar", str(OUT_RATE), "-ac", "1", "-acodec", "pcm_s16le"])
            if os.path.exists(dst) and os.path.getsize(dst) > 44:
                return
        except Exception:
            pass
        seg.set_frame_rate(OUT_RATE).set_sample_width(2).set_channels(1).export(dst, format="wav")

    def _tts_render_final(self, src, dst):
        """Дубль → готовый файл: громкость выровнена, 8 кГц."""
        self._tts_export_8k(self._tts_normalized(AudioSegment.from_file(src)), dst)

    def _tts_preview(self, src):
        """Готовая к сохранению версия дубля для прослушки (кэш во временной папке)."""
        import hashlib
        import tempfile
        st_ = os.stat(src)
        key = hashlib.md5(f"{os.path.abspath(src)}|{st_.st_mtime_ns}|{st_.st_size}".encode()).hexdigest()[:16]
        dst = os.path.join(tempfile.gettempdir(), f"gvox_tts_prev_{key}.wav")
        if not os.path.exists(dst):
            self._tts_render_final(src, dst)
        return dst

    def _tts_sample_files(self, var_key):
        """Пример значения переменной для прослушки сборки: первые готовые
        значения двух последних категорий листа (сумма — тысячи + тенге,
        дата — день + месяц)."""
        sh = self._tts()["var_sheets"].get(var_key)
        if not sh:
            return []
        files = []
        for c in sh["categories"]:
            f = next((self._tts_value_file(var_key, c["key"], v) for v in c["values"]
                      if os.path.exists(self._tts_value_file(var_key, c["key"], v))), None)
            if f:
                files.append(f)
        return files[-2:]

    def tts_chain_info(self, chain_id):
        """Кусочки цепочки с громкостью речи — чтобы на слух и на глаз было
        видно, что они из одной записи и стыкуются со значениями."""
        st = self._tts()
        ch = next((c for c in st.get("chains", []) if c["id"] == chain_id), None)
        if not ch:
            return {"error": "Нет такой фразы."}
        parts = []
        for uid in ch["parts"]:
            u = st["units"][uid]
            path = self._tts_target(u) if st.get("work_dir") else ''
            db = self._tts_speech_db(AudioSegment.from_file(path)) if path and os.path.exists(path) else None
            parts.append({"id": uid, "db": db})
        samples = {v: [os.path.basename(f) for f in self._tts_sample_files(v)] for v in ch.get("vars", [])}
        dbs = [p["db"] for p in parts if p["db"] is not None]
        return {"id": chain_id, "parts": parts, "samples": samples,
                "spread": round(max(dbs) - min(dbs), 1) if len(dbs) > 1 else 0.0}

    def tts_chain_play(self, chain_id):
        """Сборка фразы целиком: кусочки + пример значения на месте каждой
        переменной (пока значений нет — короткая пауза). Только для прослушки."""
        st = self._tts()
        ch = next((c for c in st.get("chains", []) if c["id"] == chain_id), None)
        if not ch or not st.get("work_dir"):
            return {"playing": False, "error": "Нет такой фразы."}
        segs, rate = [], None
        for k, uid in enumerate(ch["parts"]):
            u = st["units"][uid]
            if k == 0 and u.get("var_before"):
                segs.append(('var', u["var_before"]))
            path = self._tts_target(u)
            segs.append(('file', path) if os.path.exists(path) else ('gap', 600))
            if u.get("var_after"):
                segs.append(('var', u["var_after"]))
        out = None
        for kind, x in segs:
            if kind == 'file':
                parts = [AudioSegment.from_file(x)]
            elif kind == 'var':
                parts = [AudioSegment.from_file(f) for f in self._tts_sample_files(x)] or [None]
            else:
                parts = [None]
            for p in parts:
                if p is None:
                    p = AudioSegment.silent(400 if kind == 'var' else x, frame_rate=rate or 16000)
                rate = rate or p.frame_rate
                p = p.set_frame_rate(rate).set_channels(1)
                out = p if out is None else out + p
        if out is None:
            return {"playing": False}
        import tempfile
        tmp = os.path.join(tempfile.gettempdir(), f"gvox_tts_chain_{os.getpid()}.wav")
        out.export(tmp, format="wav")
        self._tts_mark_listened()
        self.player.play(tmp)
        return {"playing": True, "duration": len(out) / 1000.0}

    # ---------------- разметка таблицы: start / start_2 / end и переменные ----------------
    # Отмечается прямо во время записи (суфлёр, клавиши 1/2/3/0 и 4, 5…).
    # Отметки сразу перестраивают цепочки, хранятся в проекте и дописываются
    # в саму таблицу 1-ками (перед первой записью — резервная копия).

    def _tts_ensure_lines(self):
        st = self._tts()
        if st.get("lines") and st.get("mark_cols"):
            return True
        path = st.get("excel_path")
        if not path or not os.path.exists(path):
            return False
        parsed = parse_tts_workbook(path)
        sh = next((s for s in parsed["phrase_sheets"] if s["title"] == st.get("sheet")), None)
        if not sh:
            return False
        st.update({"lines": sh["lines"], "mark_cols": sh["mark_cols"], "var_cols": sh["var_keys"],
                   "head_row": sh["head_row"]})
        st.setdefault("marks_dirty", [])
        return True

    def _tts_rebuild_units(self):
        """Цепочки заново из строк; суфлёр остаётся на той же фразе, уже
        сохранённые файлы переезжают, если фраза стала кусочком (или наоборот)."""
        st = self._tts()
        old = {u["id"]: self._tts_target(u) for u in st["units"]} if st.get("work_dir") else {}
        anchor = None
        if hasattr(self, '_tts_prompt_items'):
            items = self._tts_prompt_items()
            pos = st.get("prompt_pos", 0)
            anchor = items[pos][-1] if 0 <= pos < len(items) else None
        st["units"], st["chains"] = build_units(st["lines"])
        for u in st["units"]:
            src = old.get(u["id"])
            if src and os.path.exists(src):
                dst = self._tts_target(u)
                if dst != src:
                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                    if os.path.exists(dst):
                        self._tts_trash(dst)
                    shutil.move(src, dst)
        if anchor is not None and hasattr(self, '_tts_prompt_items'):
            items = self._tts_prompt_items()
            st["prompt_pos"] = next((k for k, ids in enumerate(items) if anchor in ids), st.get("prompt_pos", 0))

    def tts_set_mark(self, unit_id, what, value=None):
        """what='role' — value: 'start' | 'start_2' | 'end' | None (снять);
        what='var' — value: ключ колонки (<sum_ru> → 'sum_ru'), отметка переключается."""
        st = self._tts()
        if not self._tts_ensure_lines():
            return {"error": "Не нашёл файл таблицы — разметку некуда записать."}
        if not (0 <= unit_id < len(st["lines"])):
            return {"error": "Нет такой строки."}
        ln = st["lines"][unit_id]
        if what == 'role':
            ln["role"] = value if value in FLAG_COLS else None
        elif what == 'var' and value in (st.get("var_cols") or []):
            ln["vars"] = [v for v in ln["vars"] if v != value] if value in ln["vars"] else ln["vars"] + [value]
        else:
            return {"error": "Неизвестная отметка."}
        if ln["row"] not in st.setdefault("marks_dirty", []):
            st["marks_dirty"].append(ln["row"])
        self._tts_rebuild_units()
        self._tts_save()
        threading.Thread(target=self._tts_marks_write_bg, daemon=True).start()
        out = {"state": self.tts_state()}
        if hasattr(self, 'tts_prompt_info'):
            out["prompt"] = self.tts_prompt_info()
        return out

    def _tts_marks_write_bg(self):
        lock = getattr(self, '_tts_marks_lock', None)
        if lock is None:
            lock = self._tts_marks_lock = threading.Lock()
        with lock:
            res = self.tts_marks_write()
        self._tts_push({"stage": "marks", **res})

    def tts_marks_write(self):
        """Дописать отметки в таблицу. Файл открыт в Excel — отметки ждут в проекте."""
        st = self._tts()
        dirty = list(st.get("marks_dirty") or [])
        if not dirty:
            return {"saved": True, "pending": 0}
        path = st.get("excel_path")
        if not path or not os.path.exists(path):
            return {"saved": False, "pending": len(dirty), "error": "Файл таблицы не найден."}
        backup = os.path.splitext(path)[0] + '_до_разметки.xlsx'
        try:
            if not os.path.exists(backup):
                shutil.copy2(path, backup)
            wb = openpyxl.load_workbook(path)
            ws = wb[st["sheet"]]
            cols = st.get("mark_cols") or {}
            by_row = {ln["row"]: ln for ln in st["lines"]}
            for row in dirty:
                ln = by_row.get(row)
                if not ln:
                    continue
                for f in FLAG_COLS:
                    if f in cols:
                        ws.cell(row=row, column=cols[f]).value = 1 if ln["role"] == f else None
                for v in st.get("var_cols") or []:
                    if v in cols:
                        ws.cell(row=row, column=cols[v]).value = 1 if v in ln["vars"] else None
            wb.save(path)
        except PermissionError:
            return {"saved": False, "pending": len(dirty),
                    "error": "Таблица открыта в Excel — закройте её, отметки допишутся сами."}
        except Exception as e:
            return {"saved": False, "pending": len(dirty), "error": f"Не удалось записать в таблицу: {e}"}
        st["marks_dirty"] = [r for r in st.get("marks_dirty", []) if r not in dirty]
        self._tts_save()
        return {"saved": True, "pending": len(st["marks_dirty"]), "backup": os.path.basename(backup)}
