# Демо-проект GVox для проверок и роликов: нейтральные фразы, синтетический «голос».
# Заглушки Windows-модулей (pyautogui, sounddevice, webview) — чтобы Api работал в песочнице.
import os, sys, types, shutil, wave, json, re
import numpy as np
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
for m in ('pyautogui', 'winsound', 'sounddevice'):
    sys.modules[m] = types.ModuleType(m)
wv = types.ModuleType('webview'); wv.windows = []; wv.FileDialog = types.SimpleNamespace(FOLDER=0, OPEN=1, SAVE=2)
wv.OPEN_DIALOG, wv.FOLDER_DIALOG = 1, 0
sys.modules['webview'] = wv
import openpyxl
D = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_work')   # демо-данные (в git не идут)
os.makedirs(D, exist_ok=True)
XL = os.path.join(D, 'demo.xlsx')

def make_xlsx():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Фразы'
    ws.append(['<phrase>', '<name_phrase>', '<start>', '<end>', '<sum_ru>', '<date_ru>'])
    rows = [
        ('Здравствуйте! Вас приветствует служба доставки «Город».', 'hello_ru_1', None, None, None, None),
        ('Ваш заказ передан курьеру и прибудет сегодня.', 'order_ru_1', None, None, None, None),
        ('Сумма к оплате при получении составляет', 'pay_ru_1_1', 1, None, 1, None),
        ('Оплатить можно картой или наличными.', 'pay_ru_1_2', None, 1, None, None),
        ('Курьер привезёт заказ', 'deliv_ru_2_1', 1, None, None, 1),
        ('с десяти до восемнадцати часов.', 'deliv_ru_2_2', None, 1, None, None),
        ('Стоимость доставки составит', 'cost_ru_3_1', 1, None, 1, None),
        ('Спасибо, что выбрали нас!', 'cost_ru_3_2', None, 1, None, None),
        ('Если хотите изменить время доставки, нажмите один.', 'menu_ru_1', None, None, None, None),
        ('Чтобы связаться с оператором, оставайтесь на линии.', 'oper_ru_1', None, None, None, None),
        ('Ваш звонок очень важен для нас.', 'wait_ru_1', None, None, None, None),
        ('Пожалуйста, назовите номер заказа.', 'ask_ru_1', None, None, None, None),
        ('Номер не распознан. Попробуйте ещё раз.', 'retry_ru_1', None, None, None, None),
        ('Хорошего дня! До свидания.', 'bye_ru_1', None, None, None, None),
    ]
    for r in rows:
        ws.append(list(r))
    s = wb.create_sheet('<sum_ru>')
    s.append(['<million_ru>', '<hundred_thousand_ru>', '<hundred_ru>', '<thousand_ru>', '<tenge_ru>'])
    for k in range(1, 31):
        s.append([k, (min(k, 9)) * 100, (min(k, 9)) * 100, k, k])
    d = wb.create_sheet('<date_ru>')
    d.append(['<day_ru>', '<month_ru>'])
    months = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня', 'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']
    for k in range(1, 32):
        d.append([k, months[k - 1] if k <= 12 else None])
    wb.save(XL)

R = 16000
rng = np.random.default_rng(7)

def speech(text, rate=R, level=0.25, decl=4.0):
    """Похожий на речь звук: слоги (тон+шум с огибающей), паузы между словами."""
    words = re.findall(r"[\w-]+|[.,!?]", text)
    out = [np.zeros(int(0.18 * rate))]
    for w in words:
        if w in '.,!?':
            out.append(np.zeros(int((0.28 if w in '.!?' else 0.14) * rate)))
            continue
        syl = max(1, len(re.findall(r'[аеёиоуыэюяaeiou]', w.lower())))
        for s in range(syl):
            n = int(rng.uniform(0.11, 0.19) * rate)
            t = np.arange(n) / rate
            f0 = rng.uniform(110, 170)
            tone = sum(np.sin(2 * np.pi * f0 * h * t) / h for h in (1, 2, 3, 4))
            noise = rng.standard_normal(n) * 0.35
            env = np.sin(np.pi * np.linspace(0, 1, n)) ** 1.4 * rng.uniform(0.6, 1.0)
            out.append((tone * 0.6 + noise) * env)
        out.append(np.zeros(int(rng.uniform(0.05, 0.11) * rate)))
    out.append(np.zeros(int(0.2 * rate)))
    y = np.concatenate(out)
    y = y / (np.abs(y).max() + 1e-9) * level
    y *= 10 ** (np.linspace(0, -decl, len(y)) / 20)          # к концу фразы тише
    return (y + 0.002 * rng.standard_normal(len(y))).astype(np.float32)

def write(path, y, rate=R):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with wave.open(path, 'wb') as f:
        f.setnchannels(1); f.setsampwidth(2); f.setframerate(rate)
        f.writeframes((np.clip(y, -1, 1) * 32767).astype('<i2').tobytes())

import core.tts_live as tl
from core.tts_flow import TtsFlowMixin
from core.tts_synth import TtsSynthMixin
from core.clone_prep import ClonePrepMixin

class A(TtsFlowMixin, tl.TtsLiveMixin, TtsSynthMixin, ClonePrepMixin):
    player = types.SimpleNamespace(play=lambda p: None, stop=lambda: None)

def build():
    make_xlsx()
    a = A(); a._tts_push = lambda p: None
    wd = os.path.join(D, 'Демо_проект'); shutil.rmtree(wd, ignore_errors=True)
    for sub in ('Записи', 'Чанки'):
        os.makedirs(os.path.join(wd, sub))
    a.tts_load_excel(XL); st = a._tts(); st['work_dir'] = wd
    U = st['units']
    # одна «запись» сессии: все фразы подряд, дубли — куски из неё
    raw = os.path.join(wd, 'Записи', 'Запись_демо.wav')
    parts, pos, spans = [], 0, {}
    for u in [u for u in U if u['kind'] != 'value']:
        y = speech(u['text']); spans[u['id']] = (pos, pos + len(y)); parts.append(y); pos += len(y)
    write(raw, np.concatenate(parts))
    from pydub import AudioSegment
    src = AudioSegment.from_file(raw)
    n = 0
    for uid, (p0, p1) in spans.items():
        u = U[uid]
        if u['name'] in ('retry_ru_1', 'bye_ru_1', 'ask_ru_1'):
            continue                                     # ещё не записаны
        n += 1
        a_ms, b_ms = p0 * 1000 // R, p1 * 1000 // R
        f = os.path.join('Чанки', f'{n:04d}.wav')
        src[a_ms:b_ms].export(os.path.join(wd, f), format='wav')
        t = {"file": f, "start_ms": a_ms, "ms": b_ms - a_ms, "orig": [a_ms, b_ms], "raw": raw, "heard": u['text'].lower(),
             "listened": True, "chunk": n, "unit": uid, "score": 0.96, "status": "ok"}
        st['takes'].append(t)
        if u['name'] not in ('wait_ru_1',):
            a.tts_approve(len(st['takes']) - 1)
    a._tts_save()
    return a, st, U, wd

if __name__ == '__main__':
    a, st, U, wd = build()
    print('юнитов', len(U), 'дублей', len(st['takes']), a.tts_state()['stats'])
